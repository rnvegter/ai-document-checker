import io
import sys
import unittest
import zipfile
from pathlib import Path

import docx
import pptx

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import documents  # noqa: E402
from helpers import make_docx, make_pptx  # noqa: E402


class DocxTests(unittest.TestCase):
    def build(self):
        d = docx.Document()
        d.add_heading("Report", 1)
        p = d.add_paragraph()
        p.add_run("Bold start ").bold = True
        p.add_run("plays a crucial role​.")
        d.add_table(rows=1, cols=1).cell(0, 0).text = "Cell text"
        d.sections[0].header.paragraphs[0].text = "Header line"
        d.core_properties.author = "ChatGPT"
        buf = io.BytesIO()
        d.save(buf)
        return buf.getvalue()

    def test_reads_body_tables_and_headers(self):
        doc = documents.load("x.docx", self.build())
        texts = [u.text for u in doc.units]
        self.assertEqual(texts[0], "Report")
        self.assertIn("Cell text", texts)
        self.assertIn("Header line", texts)

    def test_rewrite_keeps_paragraph_formatting_and_scrubs(self):
        doc = documents.load("x.docx", self.build())
        self.assertEqual(doc.apply({1: "It matters​.", 0: "Report", 999: "x"}), 1)
        doc.scrub("Jane Doe")
        saved = docx.Document(io.BytesIO(doc.save()))
        p = saved.paragraphs[1]
        self.assertEqual(p.text, "It matters.")
        self.assertTrue(p.runs[0].bold)
        self.assertEqual(saved.paragraphs[0].style.name, "Heading 1")
        self.assertEqual(saved.core_properties.author, "Jane Doe")
        self.assertEqual(saved.core_properties.last_modified_by, "Jane Doe")
        self.assertEqual(doc.clean_filename, "x (clean).docx")

    def test_scrub_removes_invisible_characters_everywhere(self):
        doc = documents.load("x.docx", self.build())
        doc.scrub()
        self.assertNotIn("​", documents.load("x.docx", doc.save()).review_text())

    def test_custom_properties_drop_ai_tags_but_keep_labels(self):
        doc = documents.load("x.docx", make_docx(custom=True))
        doc.scrub()
        with zipfile.ZipFile(io.BytesIO(doc.save())) as z:
            custom = z.read("docProps/custom.xml").decode()
        self.assertIn("MSIP_Label_1234_Enabled", custom)
        self.assertNotIn("GeneratedBy", custom)
        self.assertNotIn("AI_Assisted", custom)


class PptxTests(unittest.TestCase):
    def test_reads_slides_and_notes(self):
        doc = documents.load("deck.pptx", make_pptx())
        self.assertEqual(
            [(u.location, u.text) for u in doc.units],
            [
                ("Slide 1", "Quarterly plan"),
                ("Slide 1", "Moreover, synergy unlocks the full potential of teams."),
                ("Slide 1", "Budget: €40,000"),
                ("Slide 1 notes", "Mention the deadline of 1 March."),
                ("Slide 2", "Next steps"),
            ],
        )
        self.assertTrue(doc.review_text().startswith("[Slide 1]\n\nQuarterly plan"))

    def test_rewrite_and_save(self):
        doc = documents.load("deck.pptx", make_pptx())
        doc.apply({1: "The teams work better together."})
        doc.scrub("Jane")
        prs = pptx.Presentation(io.BytesIO(doc.save()))
        body = prs.slides[0].placeholders[1].text_frame
        self.assertEqual([p.text for p in body.paragraphs], ["The teams work better together.", "Budget: €40,000"])
        self.assertEqual(prs.core_properties.author, "Jane")
        self.assertEqual(prs.slides[0].notes_slide.notes_text_frame.text, "Mention the deadline of 1 March.")


class TextAndErrorTests(unittest.TestCase):
    def test_plain_text(self):
        doc = documents.load("notes.md", "One\r\n\r\nTwo\nlines\n\n\nThree".encode())
        self.assertEqual([u.text for u in doc.units], ["One", "Two\nlines", "Three"])
        doc.apply({0: "Uno"})
        self.assertEqual(doc.save().decode(), "Uno\n\nTwo\nlines\n\nThree\n")
        self.assertEqual(doc.clean_filename, "notes (clean).txt")

    def test_unsupported(self):
        for name, data in [("a.pdf", b""), ("a.doc", b""), ("a.docx", b"not a zip"), ("a.pptx", b"nope")]:
            with self.assertRaises(documents.UnsupportedDocument):
                documents.load(name, data)


if __name__ == "__main__":
    unittest.main()
