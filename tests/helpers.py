import io
import json
import zipfile
from types import SimpleNamespace

import docx
import pptx

CUSTOM_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
    'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
    '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="MSIP_Label_1234_Enabled"><vt:lpwstr>true</vt:lpwstr></property>'
    '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="3" name="GeneratedBy"><vt:lpwstr>Copilot</vt:lpwstr></property>'
    '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="4" name="AI_Assisted"><vt:lpwstr>yes</vt:lpwstr></property>'
    "</Properties>"
)


def make_docx(paragraphs=("First para.", "", "Second para."), custom=False):
    d = docx.Document()
    for text in paragraphs:
        d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    data = buf.getvalue()
    return add_custom_props(data) if custom else data


def add_custom_props(data):
    """Add a docProps/custom.xml part, as Word does for labels and tags."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for item in src.infolist():
            content = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                content = content.replace(
                    b"</Types>",
                    b'<Override PartName="/docProps/custom.xml" '
                    b'ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml"/></Types>',
                )
            elif item.filename == "_rels/.rels":
                content = content.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdCustom" Target="docProps/custom.xml" '
                    b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties"/>'
                    b"</Relationships>",
                )
            z.writestr(item, content)
        z.writestr("docProps/custom.xml", CUSTOM_XML)
    return out.getvalue()


def make_pptx():
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Quarterly plan"
    body = slide.placeholders[1].text_frame
    body.text = "Moreover, synergy unlocks the full potential of teams."
    body.add_paragraph().text = "Budget: €40,000"
    slide.notes_slide.notes_text_frame.text = "Mention the deadline of 1 March."
    second = prs.slides.add_slide(prs.slide_layouts[5])
    second.shapes.title.text = "Next steps"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def response(payload, stop_reason="end_turn"):
    text = payload if isinstance(payload, str) else json.dumps(payload)
    return SimpleNamespace(
        stop_reason=stop_reason, model="claude-opus-5-5", content=[SimpleNamespace(type="text", text=text)]
    )


class FakeMessages:
    """Stands in for client.beta.messages. `rewrite` maps a paragraph text to its new text."""

    def __init__(self, review=None, rewrite=None, stop_reason="end_turn"):
        self.review = review
        self.rewrite = rewrite or (lambda text: None)
        self.stop_reason = stop_reason
        self.calls = []

    def create(self, **kwargs):
        self.calls.append(kwargs)
        return response(self.review, self.stop_reason)

    def stream(self, **kwargs):
        self.calls.append(kwargs)
        paragraphs = json.loads(kwargs["messages"][0]["content"])
        changed = [{"id": p["id"], "text": t} for p in paragraphs if (t := self.rewrite(p["text"]))]
        changed.append({"id": 99999, "text": "unknown id is ignored"})
        final = response({"paragraphs": changed}, self.stop_reason)

        class Stream:
            def __enter__(self):
                return SimpleNamespace(get_final_message=lambda: final)

            def __exit__(self, *exc):
                return False

        return Stream()


def fake_client(**kwargs):
    return SimpleNamespace(beta=SimpleNamespace(messages=FakeMessages(**kwargs)))
