"""Read, rewrite and save .docx, .pptx and plain-text documents.

A document is a list of text units, one per paragraph. Cleaning rewrites
units in place, so Word and PowerPoint files keep their layout and styles.
"""
import io
import re
import zipfile
from collections import Counter
from dataclasses import dataclass, field

from docx import Document as open_docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree
from pptx import Presentation

from traces import scrub_text, strip_image_metadata

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

MIME = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "txt": "text/plain; charset=utf-8",
}

# Custom document properties that tag a file as AI-made. Kept narrow on purpose:
# sensitivity labels (MSIP_Label_*) and other organisational properties stay.
AI_PROPERTY = re.compile(r"copilot|chatgpt|openai|\bgpt|claude|anthropic|gemini|generated|(?<![a-z])ai(?![a-z])", re.I)


class UnsupportedDocument(ValueError):
    pass


@dataclass
class Unit:
    id: int
    text: str
    location: str = ""
    nodes: list = field(default_factory=list, repr=False)


class Document:
    def __init__(self, kind, filename, units, obj=None):
        self.kind = kind
        self.filename = filename
        self.units = units
        self._obj = obj

    @property
    def mime(self):
        return MIME[self.kind]

    @property
    def clean_filename(self):
        stem = re.sub(r"\.[^.]+$", "", self.filename) or "document"
        return f"{stem} (clean).{self.kind}"

    def review_text(self):
        """Plain text for checking; slides are marked with [Slide n]."""
        if self.kind != "pptx":
            return "\n\n".join(u.text for u in self.units)
        blocks, current = [], None
        for u in self.units:
            if u.location != current:
                current = u.location
                blocks.append(f"[{current}]")
            blocks.append(u.text)
        return "\n\n".join(blocks)

    def apply(self, rewrites):
        """Replace unit texts by id. Unknown ids and empty texts are ignored."""
        by_id = {u.id: u for u in self.units}
        changed = 0
        for uid, text in rewrites.items():
            unit = by_id.get(uid)
            text = (text or "").strip()
            if not unit or not text or text == unit.text:
                continue
            unit.text = text
            if unit.nodes:
                # The first run keeps its formatting and takes the whole text.
                first, *rest = unit.nodes
                first.text = text
                if self.kind == "docx":
                    first.set(XML_SPACE, "preserve")
                for node in rest:
                    node.text = ""
            changed += 1
        return changed

    def scrub(self, author=""):
        """Remove hidden traces and return a report of what was found.

        The report is a list of {"label", "count", "detail"} entries: hidden and
        look-alike characters in the text, AI metadata and image metadata.
        """
        counts = Counter()
        for u in self.units:
            if u.nodes:
                for node in u.nodes:
                    if node.text:
                        node.text, found = scrub_text(node.text)
                        counts.update(found)
                u.text = "".join(node.text or "" for node in u.nodes)
            else:
                u.text, found = scrub_text(u.text)
                counts.update(found)
        report = [{"label": label, "count": n, "detail": ""} for label, n in counts.items()]
        if self.kind == "txt":
            return report

        props = self._obj.core_properties
        old = [v for v in (props.author, props.last_modified_by) if v and v != author]
        if old:
            report.append({"label": "Author fields reset", "count": len(old), "detail": ", ".join(dict.fromkeys(old))})
        if props.comments:
            report.append({"label": "Document comments field cleared", "count": 1, "detail": props.comments[:80]})
        props.author = author
        props.last_modified_by = author
        props.comments = ""

        images = []
        for part in self._obj.part.package.iter_parts():
            name = str(part.partname)
            if name == "/docProps/custom.xml":
                removed = _scrub_custom_properties(part)
                if removed:
                    report.append({"label": "AI document properties", "count": len(removed), "detail": ", ".join(removed)})
            elif part.content_type.startswith("image/"):
                part._blob, removed = strip_image_metadata(part.blob)
                if removed:
                    images.append(f"{name.rsplit('/', 1)[-1]} ({', '.join(removed)})")
        if images:
            report.append({"label": "Image metadata", "count": len(images), "detail": "; ".join(images)})
        return report

    def save(self):
        if self.kind == "txt":
            return ("\n\n".join(u.text for u in self.units) + "\n").encode("utf-8")
        buf = io.BytesIO()
        self._obj.save(buf)
        return buf.getvalue()


def _scrub_custom_properties(part):
    root = part._element if hasattr(part, "_element") else etree.fromstring(part.blob)
    removed = []
    for prop in list(root):
        if AI_PROPERTY.search(prop.get("name", "")):
            removed.append(prop.get("name"))
            root.remove(prop)
    if not hasattr(part, "_element"):
        # Generic parts keep their raw bytes; python-docx/pptx expose no setter.
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    return removed


def _own_nodes(paragraph, tag, keep=lambda node: True):
    """Text nodes that belong to this paragraph, not to one nested inside it."""
    return [
        t for t in paragraph.iter(tag)
        if next(t.iterancestors(paragraph.tag), None) is paragraph and keep(t)
    ]


def _load_docx(data):
    doc = open_docx(io.BytesIO(data))
    roots = [doc.element.body]
    for rel in doc.part.rels.values():
        if rel.reltype in (RT.HEADER, RT.FOOTER) and not rel.is_external:
            roots.append(rel.target_part.element)
    units = []
    for root in roots:
        for p in root.iter(f"{W}p"):
            if any(a.tag == MC_FALLBACK for a in p.iterancestors()):
                continue  # duplicate of a text box already read from mc:Choice
            nodes = _own_nodes(p, f"{W}t")
            text = "".join(t.text or "" for t in nodes)
            if text.strip():
                units.append(Unit(len(units), text, "", nodes))
    return doc, units


def _load_pptx(data):
    prs = Presentation(io.BytesIO(data))
    units = []

    def read(root, location):
        for p in root.iter(f"{A}p"):
            # Only text runs; fields such as slide numbers stay untouched.
            nodes = _own_nodes(p, f"{A}t", lambda t: t.getparent().tag == f"{A}r")
            text = "".join(t.text or "" for t in nodes)
            if text.strip():
                units.append(Unit(len(units), text, location, nodes))

    for n, slide in enumerate(prs.slides, 1):
        read(slide._element, f"Slide {n}")
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            read(slide.notes_slide.notes_text_frame._txBody, f"Slide {n} notes")
    return prs, units


def load(filename, data):
    name = filename.lower()
    try:
        if name.endswith(".docx"):
            obj, units = _load_docx(data)
            return Document("docx", filename, units, obj)
        if name.endswith(".pptx"):
            obj, units = _load_pptx(data)
            return Document("pptx", filename, units, obj)
    except (zipfile.BadZipFile, KeyError, ValueError, etree.XMLSyntaxError) as e:
        raise UnsupportedDocument(f"That file could not be read as {name.rsplit('.', 1)[-1]}.") from e
    if name.endswith((".txt", ".md", ".markdown")):
        text = data.decode("utf-8", errors="replace").replace("\r\n", "\n")
        paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
        return Document("txt", filename, [Unit(i, p) for i, p in enumerate(paras)])
    if name.endswith((".doc", ".ppt")):
        raise UnsupportedDocument("Old .doc/.ppt files are not supported. Save as .docx or .pptx first.")
    raise UnsupportedDocument("Upload a .docx, .pptx, .txt or .md file.")
